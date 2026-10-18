import re
from collections.abc import Iterable
from pathlib import Path

from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.core.access import ACCESS_LEVELS
from app.core.config import BASE_DIR

SUPPORTED = {".pdf", ".txt", ".md", ".docx"}

_markdown_sections = MarkdownHeaderTextSplitter(
    headers_to_split_on=[("#", "title"), ("##", "section")], strip_headers=True
)
_splitter = RecursiveCharacterTextSplitter(chunk_size=900, chunk_overlap=120, add_start_index=True)


def source_id(path: Path) -> str:
    """Stable, machine-independent source name: project-relative when possible."""
    path = path.resolve()
    try:
        return path.relative_to(BASE_DIR).as_posix()
    except ValueError:
        # Outside the project (e.g. a custom UPLOAD_DIR): keep the department folder.
        return "/".join(path.parts[-2:])


def _title_from_name(path: Path) -> str:
    return path.stem.replace("_", " ").replace("-", " ").strip().title()


def load_file(path: Path) -> list[Document]:
    suffix = path.suffix.lower()
    source = source_id(path)
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return [
            Document(page_content=text, metadata={"source": source, "page": i + 1, "title": _title_from_name(path)})
            for i, page in enumerate(reader.pages)
            if (text := (page.extract_text() or "").strip())
        ]
    if suffix == ".md":
        return [Document(page_content=path.read_text(encoding="utf-8"), metadata={"source": source, "format": "markdown"})]
    if suffix == ".txt":
        return [Document(page_content=path.read_text(encoding="utf-8"), metadata={"source": source, "title": _title_from_name(path)})]
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": source, "title": _title_from_name(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def _sections(doc: Document) -> list[Document]:
    """Split Markdown by headings so every section keeps its document title and heading."""
    if doc.metadata.get("format") != "markdown":
        return [doc]
    base = {k: v for k, v in doc.metadata.items() if k != "format"}
    parts = [p for p in _markdown_sections.split_text(doc.page_content) if p.page_content.strip()]
    title = next((p.metadata["title"] for p in parts if p.metadata.get("title")), None) or _title_from_name(Path(base["source"]))
    # A short preamble before the first "##" (owner, effective date, status) describes the whole
    # document: attach it to every section instead of indexing it as a near-empty chunk.
    if len(parts) > 1 and not parts[0].metadata.get("section") and len(parts[0].page_content) < 300:
        base["doc_info"] = " ".join(parts.pop(0).page_content.split())
    return [
        Document(page_content=p.page_content, metadata={**base, "title": title, **({"section": p.metadata["section"]} if p.metadata.get("section") else {})})
        for p in parts
    ]


_STATUS_RE = re.compile(r"\bStatus:\s*(SUPERSEDED|ARCHIVED|CURRENT)", re.IGNORECASE)
_ACCESS_RE = re.compile(r"\bAccess:\s*([a-z_]+)", re.IGNORECASE)


def lifecycle_metadata(doc_info: str, access_override: str | None = None) -> dict:
    """Read document status and access level from the document header.

    Defaults are deliberately safe for retrieval filters: a document is "current" unless marked
    SUPERSEDED/ARCHIVED, and readable by "all" unless an Access level is declared.
    """
    status = _STATUS_RE.search(doc_info or "")
    access = _ACCESS_RE.search(doc_info or "")
    level = (access_override or (access.group(1) if access else "all")).lower()
    if level not in ACCESS_LEVELS:
        raise ValueError(f"Unknown access level '{level}'. Valid: {', '.join(ACCESS_LEVELS)}")
    return {
        "status": "superseded" if status and status.group(1).upper() in ("SUPERSEDED", "ARCHIVED") else "current",
        "access": level,
    }


def chunk_documents(docs: Iterable[Document], access_override: str | None = None) -> list[Document]:
    """Section-aware chunking with a contextual header on every chunk.

    Each chunk starts with "<document title> > <section>" so that both the embedding and
    the reranker see what the chunk is about, even when the body text never repeats it.
    Every chunk also gets `status` and `access` metadata for retrieval-time filtering.
    """
    chunks = []
    for doc in docs:
        for section in _sections(doc):
            section.metadata.update(lifecycle_metadata(section.metadata.get("doc_info", ""), access_override))
            for chunk in _splitter.split_documents([section]):
                header = " > ".join(x for x in (chunk.metadata.get("title"), chunk.metadata.get("section")) if x)
                if chunk.metadata.get("doc_info"):
                    header += f" [{chunk.metadata['doc_info']}]"
                if header:
                    chunk.page_content = f"{header}\n{chunk.page_content}"
                chunks.append(chunk)
    return chunks
